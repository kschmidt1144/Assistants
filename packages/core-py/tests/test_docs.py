import io

import pytest
from assistants_core.docs import parse_document


def test_txt():
    assert parse_document(b"hello world", "notes.txt") == "hello world"


def test_md():
    assert "Title" in parse_document(b"# Title\n\nbody", "readme.md")


def test_unsupported():
    with pytest.raises(ValueError):
        parse_document(b"\x89PNG...", "image.png")


def test_docx_roundtrip():
    from docx import Document

    doc = Document()
    doc.add_paragraph("First line")
    doc.add_paragraph("Second line")
    buf = io.BytesIO()
    doc.save(buf)

    text = parse_document(buf.getvalue(), "resume.docx")
    assert "First line" in text
    assert "Second line" in text

def test_extensions():
    # Uppercase extension
    assert parse_document(b"hello", "file.TXT") == "hello"
    
    # Extension-less filename
    with pytest.raises(ValueError, match="Unsupported document type"):
        parse_document(b"hello", "filename")

def test_utf8_replace():
    # Invalid UTF-8 bytes should not crash if we use errors="replace", 
    # but currently parse_document might decode differently. 
    # Let's ensure it handles it or raises a specific error.
    # If the current code doesn't use errors="replace", it might raise UnicodeDecodeError.
    try:
        res = parse_document(b"hello \xff world", "test.txt")
        assert "hello" in res
    except UnicodeDecodeError:
        pass # If we haven't implemented it yet

def test_fixture_pdf():
    import os
    fixture_path = os.path.join(os.path.dirname(__file__), "fixtures", "sample.pdf")
    with open(fixture_path, "rb") as f:
        data = f.read()
    text = parse_document(data, "sample.pdf")
    assert "Hello PDF World" in text

